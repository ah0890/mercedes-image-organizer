"""Segregate Mercedes_Organized/_Remaining/Not Matched/ by CLASS NAME only.

    python segregate_not_matched.py --dry-run     # report only, move nothing
    python segregate_not_matched.py               # move files into Not Matched/<Class>/

Scope: ONLY files inside Not Matched/ (directly, or in a class subfolder this
script created) are read or moved. Nothing else in the project is touched.

Class names are read from the Markdown page list (every "## <Class> - N pages"
heading that has page bullets). For each class the accepted spellings are
generated, never hard-coded:

    "A-Class" -> A-Class, A Class, A_Class, aclass, Class A, Class-A, Class_A (any case)
    "AMG GT"  -> AMG GT, AMG-GT, AMG_GT, amggt                                (any case)
    "GLE"     -> GLE (any case) as a whole word

Matching is token-boundary aware: "a" inside a word, "amg" inside "amgx", or
"sl" inside "sls" never counts. Model numbers and chassis codes (A200, W176 ...)
are NOT used here.

A file is moved only when exactly ONE class matches. It stays directly in
Not Matched/ when:
  * no class name is found                       -> NOT MATCHED
  * more than one class name is found            -> AMBIGUOUS
  * its name starts with another model family    -> AMBIGUOUS
    ("mercedes-benz-sls amg-... sls amg gt": the file's own family is SLS AMG,
     "amg gt" only appears later in the name)

Nothing is overwritten or deleted: an identical file already at the destination
is reported as DUPLICATE and left in place; a different file with the same name
is moved under a new name ("name (2).webp") and reported as NAME CONFLICT.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from mercedes_organizer.parser import parse_markdown
from mercedes_organizer.utils import fs_path, safe_folder_name, strip_accents

PROJECT = Path(__file__).resolve().parent
NOT_MATCHED = PROJECT / "Mercedes_Organized" / "_Remaining" / "Not Matched"
MARKDOWN = PROJECT / "EF_Mercedes_Full_Page_List.md"
REPORT_NAME = "class_name_segmentation_report.xlsx"

SEP = r"[\s_\-]"          # space, underscore or hyphen
BRAND = re.compile(r"^mercedes[\s_\-]+benz[\s_\-]+")

MOVED, NOT_MATCHED_ST, AMBIGUOUS, DUPLICATE, CONFLICT = (
    "MOVED", "NOT MATCHED", "AMBIGUOUS", "DUPLICATE", "MOVED (NAME CONFLICT - RENAMED)")


def norm(text: str) -> str:
    return strip_accents(text).lower().replace("–", "-").replace("—", "-")


def class_patterns(class_name: str) -> list[re.Pattern]:
    """Spelling variants of one class name, generated from the name itself."""
    words = [w for w in re.split(r"[\s_\-]+", norm(class_name)) if w]
    b, e = r"(?<![a-z0-9])", r"(?![a-z0-9])"
    if len(words) == 2 and "class" in words:
        letter = words[0] if words[1] == "class" else words[1]
        x = re.escape(letter)
        return [re.compile(b + x + SEP + "*class" + e),          # A-Class, A Class, A_Class, aclass
                re.compile(b + "class" + SEP + "+" + x + e)]     # Class A, Class-A, Class_A
    return [re.compile(b + (SEP + "*").join(re.escape(w) for w in words) + e)]   # AMG GT, AMG-GT, amggt, GLE


def leading_family(name: str) -> str | None:
    """Family segment of a 'mercedes-benz-<family>-...' file name, else None."""
    m = BRAND.match(name)
    if not m:
        return None
    rest = name[m.end():]
    cls = re.match(r"[a-z]+-class(?![a-z0-9])", rest)
    return cls.group(0) if cls else rest.split("-", 1)[0].strip() or None


@dataclass
class Row:
    location: str
    file: str
    status: str
    matched: str = ""
    pattern: str = ""
    destination: str = ""
    possible: tuple[str, ...] = ()
    reason: str = ""


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(fs_path(path), "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def classify(filename: str, patterns: dict[str, list[re.Pattern]]) -> tuple[list[tuple[str, str]], str | None]:
    name = norm(Path(filename).stem)
    found = []
    for cls, pats in patterns.items():
        for p in pats:
            m = p.search(name)
            if m:
                found.append((cls, filename[m.start():m.end()] if len(filename) >= m.end() else m.group(0)))
                break
    return found, leading_family(name)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--dry-run", action="store_true", help="report only, move nothing")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    if not NOT_MATCHED.is_dir():
        print(f"ERROR: {NOT_MATCHED} does not exist", file=sys.stderr)
        return 2

    classes = [c.name for c in parse_markdown(MARKDOWN)]
    patterns = {c: class_patterns(c) for c in classes}
    folder_of = {c: safe_folder_name(c) for c in classes}
    class_folders = set(folder_of.values())

    # Files to check: directly in Not Matched, plus files already in one of OUR class subfolders.
    todo: list[tuple[Path, str | None]] = []
    for entry in sorted(os.listdir(fs_path(NOT_MATCHED))):
        full = NOT_MATCHED / entry
        if os.path.isfile(fs_path(full)) and entry != REPORT_NAME and not entry.startswith("~$"):
            todo.append((full, None))
        elif os.path.isdir(fs_path(full)) and entry in class_folders:
            for sub in sorted(os.listdir(fs_path(full))):
                if os.path.isfile(fs_path(full / sub)):
                    todo.append((full / sub, entry))

    rows: list[Row] = []
    created: list[str] = []
    existing_dirs = {e for e in os.listdir(fs_path(NOT_MATCHED)) if os.path.isdir(fs_path(NOT_MATCHED / e))}
    for path, already_in in todo:
        location = "Not Matched/" + (f"{already_in}/" if already_in else "")
        found, family = classify(path.name, patterns)
        cls_names = [c for c, _ in found]
        if not found:
            rows.append(Row(location, path.name, NOT_MATCHED_ST, destination="Remaining in Not Matched/",
                            reason="no class name from the page list in the file name"))
            continue
        if len(found) > 1:
            rows.append(Row(location, path.name, AMBIGUOUS, possible=tuple(cls_names),
                            reason="file name contains more than one class name: " + ", ".join(
                                f"{c} ('{p}')" for c, p in found), destination="Remaining in Not Matched/"))
            continue
        cls, pat = found[0]
        if family is not None and not any(p.fullmatch(family) for p in patterns[cls]):
            rows.append(Row(location, path.name, AMBIGUOUS, possible=(cls, family.upper() if len(family) <= 3
                                                                      else family.title()),
                            reason=f"file name starts with the '{family}' family; '{pat}' only appears later",
                            destination="Remaining in Not Matched/"))
            continue
        folder = folder_of[cls]
        if already_in == folder:
            rows.append(Row(location, path.name, "ALREADY IN PLACE", cls, pat, f"Not Matched/{folder}/"))
            continue
        dest_dir = NOT_MATCHED / folder
        dest = dest_dir / path.name
        status = MOVED
        if os.path.exists(fs_path(dest)):
            if sha256(dest) == sha256(path):
                rows.append(Row(location, path.name, DUPLICATE, cls, pat, f"Not Matched/{folder}/",
                                reason="identical file already in the class folder - left in place"))
                continue
            stem, n = path.stem, 2
            while os.path.exists(fs_path(dest)):
                dest = dest_dir / f"{stem} ({n}){path.suffix}"
                n += 1
            status = CONFLICT
        if not args.dry_run:
            if folder not in existing_dirs:
                os.makedirs(fs_path(dest_dir), exist_ok=True)
                existing_dirs.add(folder)
                created.append(folder)
            os.replace(fs_path(path), fs_path(dest))
        elif folder not in existing_dirs and folder not in created:
            created.append(folder)
        rows.append(Row(location, path.name, status, cls, pat, f"Not Matched/{folder}/" +
                        (dest.name if status == CONFLICT else "")))

    write_report(rows, classes, created, args.dry_run)
    print_summary(rows, classes, created, args.dry_run)
    return 0


# ----------------------------------------------------------------------------- report

HEADER = PatternFill("solid", start_color="1F3864")
FILLS = {MOVED: "E2F0D9", CONFLICT: "E2F0D9", "ALREADY IN PLACE": "E2F0D9", AMBIGUOUS: "FFF2CC",
         NOT_MATCHED_ST: "F2F2F2", DUPLICATE: "F2F2F2"}


def _sheet(ws, headers, rows, widths, status_col=None):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
        if status_col is not None:
            cell = ws.cell(ws.max_row, status_col + 1)
            if cell.value in FILLS:
                cell.fill = PatternFill("solid", start_color=FILLS[cell.value])
                cell.font = Font(bold=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def write_report(rows: list[Row], classes: list[str], created: list[str], dry_run: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    moved = [r for r in rows if r.status in (MOVED, CONFLICT)]
    per_class = Counter(r.matched for r in rows if r.status in (MOVED, CONFLICT, "ALREADY IN PLACE"))
    summary = [
        ["Mode", "DRY RUN - nothing moved" if dry_run else "Files moved"],
        ["Total Files Checked", len(rows)],
        ["Files Matched", sum(r.status in (MOVED, CONFLICT, DUPLICATE, "ALREADY IN PLACE") for r in rows)],
        ["Files Moved", len(moved)],
        ["Files Not Matched", sum(r.status == NOT_MATCHED_ST for r in rows)],
        ["Ambiguous Files", sum(r.status == AMBIGUOUS for r in rows)],
        ["Duplicates (left in place)", sum(r.status == DUPLICATE for r in rows)],
        ["Name conflicts (moved, renamed)", sum(r.status == CONFLICT for r in rows)],
        ["Classes Detected", f"{len(classes)} (from {MARKDOWN.name})"],
        ["Class Folders Created", f"{len(created)}" + (f": {', '.join(created)}" if created else "")],
    ]
    _sheet(ws, ["Item", "Value"], summary, [32, 90])
    ws.append([])
    ws.append(["Class", "Files Assigned"])
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER
    for cls in classes:
        if per_class[cls]:
            ws.append([cls, per_class[cls]])
    ws.append(["Classes with no files", ", ".join(c for c in classes if not per_class[c]) or "—"])

    ws = wb.create_sheet("Class Segmentation")
    _sheet(ws, ["Original Location", "File Name", "Matched Class", "Matched Pattern", "Status", "Destination"],
           [[r.location, r.file, r.matched or "—", r.pattern or "—", r.status, r.destination or "—"]
            for r in sorted(rows, key=lambda r: (r.status != MOVED, r.matched, r.file))],
           [22, 110, 14, 18, 22, 34], status_col=4)

    ws = wb.create_sheet("Ambiguous")
    amb = [r for r in rows if r.status == AMBIGUOUS]
    if amb:
        _sheet(ws, ["File Name", "Possible Class 1", "Possible Class 2", "Reason", "Action"],
               [[r.file, r.possible[0], r.possible[1] if len(r.possible) > 1 else "—", r.reason,
                 "LEFT IN NOT MATCHED"] for r in amb], [110, 18, 18, 70, 22])
    else:
        ws["A1"] = "NO AMBIGUOUS FILES"

    path = NOT_MATCHED / REPORT_NAME
    try:
        wb.save(fs_path(path))
    except PermissionError:
        import datetime as dt
        path = path.with_name(f"{path.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        wb.save(fs_path(path))
        print(f"NOTE: report was open in Excel - saved as {path.name}")


def print_summary(rows: list[Row], classes: list[str], created: list[str], dry_run: bool) -> None:
    per_class = Counter(r.matched for r in rows if r.status in (MOVED, CONFLICT, "ALREADY IN PLACE"))
    bar = "=" * 40
    lines = [bar, "NOT MATCHED - CLASS NAME SEGREGATION" + (" (DRY RUN)" if dry_run else ""), bar, "",
             "Source:", "Remaining/Not Matched/", "",
             "Classes Detected from MD:", str(len(classes)), "",
             "Files Checked:", str(len(rows)), "",
             "Files Segregated:", str(sum(per_class.values())), "",
             "Files Still Not Matched:", str(sum(r.status == NOT_MATCHED_ST for r in rows)), "",
             "Ambiguous Files:", str(sum(r.status == AMBIGUOUS for r in rows)), "",
             "Class Folders Created:", str(len(created)), "", "-" * 40, ""]
    lines += [f"{cls}: {per_class[cls]} files" for cls in classes if per_class[cls]]
    lines += ["", "-" * 40, "", "Unmatched and ambiguous files remain in:", "", "Remaining/Not Matched/", "",
              "Excel Report:", "", f"Remaining/Not Matched/{REPORT_NAME}", "", bar]
    print("\n".join(lines))


if __name__ == "__main__":
    sys.exit(main())
