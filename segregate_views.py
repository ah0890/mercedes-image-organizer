"""Second-level sort INSIDE each class folder, by the view named in the file name.

    python segregate_views.py --dry-run            # plan + report only, move nothing
    python segregate_views.py                      # move matching images into <Class>/<View>/
    python segregate_views.py --also "Back View"   # additionally switch on a defined category

Scope
-----
* Class folders are the top-level folders of Mercedes_Organized/ that do not
  start with "_" - discovered on disk, never hard-coded.
* Only image files directly in a class folder are moved, and only into a view
  subfolder of the SAME class folder. _Remaining, _Duplicates, _Unrelated Files,
  the source folder and the class organisation itself are never touched.
* Categories come from mercedes_organizer/view_categories.py. Only categories
  that are enabled there (now: "45 Angle Front View") or passed with --also are used.

Safety
------
* The planned moves are written to view_category_moves.csv (with SHA-256)
  BEFORE anything moves; that file is also the undo record.
* Never overwrites: if the destination name exists and the content is
  identical the file is left in place (reported); if it differs, the file is
  moved under a collision-safe name "name (2).webp" (reported).
* Re-running is safe: files already inside a category subfolder are reported
  as ALREADY ORGANIZED and not moved; folders are reused, never duplicated.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import logging
import os
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from mercedes_organizer.catalog import sha256_of
from mercedes_organizer.utils import fs_path
from mercedes_organizer.view_categories import ViewCategory, detect, enabled_categories, VIEW_CATEGORIES

PROJECT = Path(__file__).resolve().parent
OUTPUT = PROJECT / "Mercedes_Organized"
REPORT = PROJECT / "view_category_segmentation_report.xlsx"
LOG = PROJECT / "view_category_segmentation.log"
PLAN = PROJECT / "view_category_moves.csv"
IMAGE_EXTS = {".webp", ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".avif"}

MOVED, PLANNED, ALREADY, NOT_MATCHED, SKIP_IDENTICAL, ERROR = (
    "MOVED", "PLANNED (dry run)", "ALREADY ORGANIZED", "NOT MATCHED", "SKIPPED - identical file at destination",
    "ERROR")

log = logging.getLogger("view_segmentation")


@dataclass
class Item:
    class_folder: str
    path: Path                       # current location
    category: str = ""
    pattern: str = ""
    destination: str = ""
    action: str = ""
    status: str = ""
    collision: str = "No"
    notes: str = ""
    sha256: str = field(default="", repr=False)
    original: str = ""               # where the file was before an earlier run moved it


def class_folders(root: Path) -> list[Path]:
    return sorted(p for p in root.iterdir() if p.is_dir() and not p.name.startswith("_"))


def move_history() -> dict[str, tuple[str, str]]:
    """Destination (as written in the undo record) -> (original path, run time) of earlier real moves."""
    history: dict[str, tuple[str, str]] = {}
    if PLAN.exists():
        for r in csv.DictReader(open(PLAN, encoding="utf-8")):
            if r.get("mode") == "move":
                history[os.path.normcase(r["to"].replace("\\", "/"))] = (r["from"], r["run"])
    return history


def plan(root: Path, categories: list[ViewCategory]) -> list[Item]:
    known = {c.folder for c in VIEW_CATEGORIES}
    by_folder = {c.folder: c for c in VIEW_CATEGORIES}
    history = move_history()
    items: list[Item] = []
    for cls in class_folders(root):
        for entry in sorted(os.listdir(fs_path(cls))):
            full = cls / entry
            if os.path.isdir(fs_path(full)):
                if entry in known:                                   # a category folder we manage
                    for name in sorted(os.listdir(fs_path(full))):
                        if Path(name).suffix.lower() in IMAGE_EXTS:
                            hit = detect(name, [by_folder[entry]])
                            item = Item(cls.name, full / name, entry, hit[1] if hit else "", f"{cls.name}/{entry}/",
                                        "none", ALREADY, notes="already inside its view subfolder")
                            past = history.get(os.path.normcase(f"{cls.name}/{entry}/{name}"))
                            if past:
                                item.original = past[0]
                                item.notes = f"moved here on {past[1]}; already inside its view subfolder"
                            items.append(item)
                else:
                    log.warning("Unknown subfolder left untouched: %s", full)
                continue
            if Path(entry).suffix.lower() not in IMAGE_EXTS:
                continue
            hit = detect(entry, categories)
            if not hit:
                items.append(Item(cls.name, full, action="none", status=NOT_MATCHED,
                                  notes="file name does not contain an enabled view category"))
                continue
            cat, text = hit
            dest_dir = cls / cat.folder
            dest = dest_dir / entry
            item = Item(cls.name, full, cat.folder, text, f"{cls.name}/{cat.folder}/", "move", PLANNED,
                        sha256=sha256_of(full))
            if os.path.exists(fs_path(dest)):
                if sha256_of(dest) == item.sha256:
                    item.action, item.status, item.collision = "none", SKIP_IDENTICAL, "Yes (identical)"
                    item.notes = "same file already in the view subfolder; left in place"
                else:
                    stem, n = Path(entry).stem, 2
                    while os.path.exists(fs_path(dest)):
                        dest = dest_dir / f"{stem} ({n}){Path(entry).suffix}"
                        n += 1
                    item.collision = "Yes (renamed)"
                    item.notes = f"different file with the same name exists; will be saved as {dest.name}"
            item.destination = f"{cls.name}/{cat.folder}/{dest.name}"
            items.append(item)
    return items


def execute(root: Path, items: list[Item], dry_run: bool) -> None:
    for it in items:
        if it.action != "move":
            continue
        if dry_run:
            continue
        dest = root / it.destination
        try:
            os.makedirs(fs_path(dest.parent), exist_ok=True)
            if os.path.exists(fs_path(dest)):                  # appeared after planning: never overwrite
                raise FileExistsError(f"destination exists: {dest}")
            os.replace(fs_path(it.path), fs_path(dest))
            if sha256_of(dest) != it.sha256:
                raise OSError("hash changed after move")
            it.status = MOVED
            log.info("MOVED %s -> %s", it.path.relative_to(root), it.destination)
        except OSError as exc:
            it.status, it.notes = ERROR, str(exc)
            log.error("Move failed %s: %s", it.path, exc)


def write_plan(items: list[Item], root: Path, dry_run: bool) -> None:
    """Append this run's planned moves (the undo record is kept across runs)."""
    new = not PLAN.exists()
    run = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(PLAN, "a", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        if new:
            w.writerow(["run", "mode", "class_folder", "from", "to", "sha256", "collision"])
        for it in items:
            if it.action == "move":
                w.writerow([run, "dry-run" if dry_run else "move", it.class_folder,
                            str(it.path.relative_to(root)), it.destination, it.sha256, it.collision])


# ----------------------------------------------------------------------------- report

_HDR = PatternFill("solid", start_color="1F3864")
_FILL = {MOVED: "E2F0D9", PLANNED: "DDEBF7", ALREADY: "E2F0D9", NOT_MATCHED: "F2F2F2", SKIP_IDENTICAL: "FFF2CC",
         ERROR: "FCE4E4"}


def _table(ws, headers, rows, widths, status_col=None):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    first = ws.max_row
    for r in rows:
        ws.append(r)
        if status_col is not None and r[status_col] in _FILL:
            cell = ws.cell(ws.max_row, status_col + 1)
            cell.fill, cell.font = PatternFill("solid", start_color=_FILL[r[status_col]]), Font(bold=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = f"A{first + 1}"
    if rows:
        ws.auto_filter.ref = f"A{first}:{chr(64 + len(headers))}{ws.max_row}"


def summary(items: list[Item], categories: list[ViewCategory], folders: list[str]) -> dict:
    matched = [i for i in items if i.category]
    return {
        "folders": len(folders),
        "scanned": len(items),
        "matched": Counter(i.category for i in matched),
        "moved": Counter(i.category for i in items if i.status in (MOVED, PLANNED)),
        "already": Counter(i.category for i in items if i.status == ALREADY),
        "not_matched": sum(i.status == NOT_MATCHED for i in items),
        "collisions": sum(i.collision != "No" for i in items),
        "errors": sum(i.status == ERROR for i in items),
        "per_class": {f: sum(1 for i in matched if i.class_folder == f) for f in folders},
        "categories": [c.folder for c in categories],
        "moved_all_runs": sum(1 for i in items if i.original) + sum(i.status == MOVED for i in items),
    }


def write_report(items: list[Item], s: dict, root: Path, dry_run: bool) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Movement Report"
    rows = [[i.class_folder, i.original or str(i.path.relative_to(root)), i.path.name, i.category or "—", i.pattern or "—",
             (i.destination.rsplit("/", 1)[0] + "/") if i.category else "—", i.action, i.status, i.collision, i.notes]
            for i in items if i.category]
    _table(ws, ["Original Class Folder", "Original File Path", "File Name", "Detected Category", "Matched Pattern",
                "Destination Folder", "Action", "Status", "Collision", "Notes"],
           rows, [16, 70, 60, 20, 22, 30, 9, 22, 14, 40], status_col=7)

    ws = wb.create_sheet("Summary")
    cat_lines = []
    for cat in s["categories"]:
        cat_lines += [[f"Total files matching {cat}", s["matched"][cat]],
                      [f"Total files {'planned to move' if dry_run else 'moved'} - {cat}", s["moved"][cat]],
                      [f"Total files already organized - {cat}", s["already"][cat]]]
    info = [
        ["Generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S") + ("   (DRY RUN - nothing moved)" if dry_run else "")],
        ["Enabled categories", ", ".join(s["categories"])],
        ["Total class folders scanned", s["folders"]],
        ["Total image files scanned", s["scanned"]],
        *cat_lines,
        ["Total files moved (all categories, this run)", sum(s["moved"].values()) if not dry_run else 0],
        ["Total files moved (all runs, from the undo record)", s["moved_all_runs"]],
        ["Total files already organized", sum(s["already"].values())],
        ["Total files not matching", s["not_matched"]],
        ["Total filename collisions", s["collisions"]],
        ["Total errors", s["errors"]],
    ]
    _table(ws, ["Item", "Value"], info, [44, 60])
    ws.append([])
    _table(ws, ["Class Folder", "Files in view subfolders"], [[f, n] for f, n in s["per_class"].items()], [44, 60])
    ws.freeze_panes = None

    ws = wb.create_sheet("Unmatched - Not Moved")
    rows = [[i.class_folder, i.path.name, i.notes, "left in class folder"] for i in items
            if i.status in (NOT_MATCHED, SKIP_IDENTICAL, ERROR)]
    _table(ws, ["Class Folder", "File Name", "Reason", "Action"], rows, [16, 110, 55, 22])

    path = REPORT
    try:
        wb.save(path)
    except PermissionError:
        path = REPORT.with_name(f"{REPORT.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        wb.save(path)
        log.error("%s was open in Excel - saved as %s", REPORT.name, path.name)
    return path


def main() -> int:
    ap = argparse.ArgumentParser(description="Sort images inside each class folder by view category.")
    ap.add_argument("--dry-run", action="store_true", help="plan and report only")
    ap.add_argument("--also", action="append", default=[], metavar="CATEGORY",
                    help="also enable a defined category, e.g. --also \"Back View\"")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")

    handler = logging.FileHandler(LOG, mode="a", encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s"))
    log.addHandler(handler)
    log.setLevel(logging.INFO)
    started = dt.datetime.now()
    log.info("=== START view category segmentation (dry_run=%s) ===", args.dry_run)

    categories = enabled_categories(args.also)
    unknown = [a for a in args.also if a.lower() not in {c.folder.lower() for c in VIEW_CATEGORIES}]
    if unknown:
        print(f"ERROR: unknown category {unknown}; defined: {[c.folder for c in VIEW_CATEGORIES]}", file=sys.stderr)
        return 2
    folders = [f.name for f in class_folders(OUTPUT)]
    log.info("Categories: %s | class folders scanned (%d): %s", [c.folder for c in categories], len(folders), folders)

    items = plan(OUTPUT, categories)
    write_plan(items, OUTPUT, args.dry_run)
    planned = sum(i.action == "move" for i in items)
    log.info("Plan: %d file(s) to move, written to %s", planned, PLAN.name)
    print(f"Planned moves: {planned} (listed in {PLAN.name})")

    execute(OUTPUT, items, args.dry_run)
    s = summary(items, categories, folders)
    report = write_report(items, s, OUTPUT, args.dry_run)

    moved = sum(i.status == MOVED for i in items)
    log.info("Files scanned %d | moved %d | already organized %d | not matched %d | collisions %d | errors %d",
             s["scanned"], moved, sum(s["already"].values()), s["not_matched"], s["collisions"], s["errors"])
    log.info("=== END (%.1fs) ===", (dt.datetime.now() - started).total_seconds())

    lines = ["", "Mercedes View Category Segmentation Complete" + (" (DRY RUN - nothing moved)" if args.dry_run else ""),
             "", f"Class folders scanned: {s['folders']}", f"Image files scanned: {s['scanned']}", ""]
    for cat in s["categories"]:
        lines += [f"{cat}:", f"    Matched: {s['matched'][cat]}",
                  f"    {'Would move' if args.dry_run else 'Moved'}: {s['moved'][cat]}",
                  f"    Already organized: {s['already'][cat]}", ""]
    lines += [f"Not matched: {s['not_matched']}", f"Collisions: {s['collisions']}", f"Errors: {s['errors']}", "",
              "Class-wise results:", *[f"{f}: {n}" for f, n in s["per_class"].items()], "",
              "Report:", report.name, "", "Log:", LOG.name, ""]
    print("\n".join(lines))
    return 1 if s["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
